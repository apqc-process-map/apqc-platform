"""Validate a private Process Map workbook and emit reviewable SQL batches.

The workbook and generated payload stay outside Git. Only this importer and schema are versioned.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

APQC_SHEET_PREFIX = "APQC_CAT_ver2"
DOCUMENT_SHEET_PREFIX = "Process_MAP_VBC"
SOURCE_NAME = "Process Map.xlsx / Process_MAP_VBC1009"
NAMESPACE = uuid.UUID("a9f62f0f-6407-4dde-9b9f-c34310b07d9d")
TYPE_MAP = {
    "thông báo": ("announcement", "communication"),
    "quy trình": ("procedure", "execution"),
    "quy định": ("regulation", "governance"),
    "quyết định": ("decision", "governance"),
    "hướng dẫn": ("work_instruction", "execution"),
    "chính sách": ("policy", "governance"),
    "tiêu chuẩn": ("standard", "governance"),
    "quy chế": ("regulation", "governance"),
    "nội quy": ("regulation", "governance"),
    "thỏa ước lao động tập thể": ("contract", "external"),
}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalized_label(value: object) -> str:
    return re.sub(r"\s+", " ", text(value)).casefold()


def iso_date(value: object) -> str | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = text(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    return None


def sheet_with_prefix(workbook: openpyxl.Workbook, prefix: str) -> str:
    matches = [name for name in workbook.sheetnames if name.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one sheet starting with {prefix!r}; found {matches}")
    return matches[0]


def row_dicts(sheet, header_row: int, first_data_row: int):
    headers = [text(value) for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    for row_number, values in enumerate(sheet.iter_rows(min_row=first_data_row, values_only=True), start=first_data_row):
        yield row_number, {header: value for header, value in zip(headers, values) if header}


def sql_text(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"

def sql_json(rows: list[dict], tag: str) -> str:
    payload = json.dumps(rows, ensure_ascii=False, separators=(",", ":"), default=str)
    if f"${tag}$" in payload:
        raise ValueError("Unexpected SQL dollar-quote marker in source data")
    return f"${tag}${payload}${tag}$::jsonb"


def write(path: Path, content: str) -> None:
    path.write_text(content.rstrip() + "\n", encoding="utf-8", newline="\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--document-batch-size", type=int, default=200)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    workbook_bytes = args.workbook.read_bytes()
    checksum = hashlib.sha256(workbook_bytes).hexdigest()
    workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    apqc_sheet_name = sheet_with_prefix(workbook, APQC_SHEET_PREFIX)
    document_sheet_name = sheet_with_prefix(workbook, DOCUMENT_SHEET_PREFIX)

    apqc_rows: list[dict] = []
    for order, (row_number, row) in enumerate(row_dicts(workbook[apqc_sheet_name], 6, 7), start=1):
        code = text(row.get("Mã L3 APQC"))
        if not code:
            continue
        apqc_rows.append({
            "source_row": row_number,
            "display_order": order,
            "category_code": text(row.get("Category")),
            "parent_code": text(row.get("L2 cha")),
            "hierarchy_id": code,
            "pcf_id": text(row.get("Element ID")),
            "name": text(row.get("Tên Process APQC (EN)")),
            "description": text(row.get("Original description (EN)")) or None,
            "pnj_description": text(row.get("Diễn giải PNJ (đề xuất)")) or None,
            "trigger_text": text(row.get("Trigger")) or None,
            "output_text": text(row.get("Output")) or None,
            "select_when": text(row.get("Chọn khi")) or None,
            "do_not_select_when": text(row.get("Không chọn khi")) or None,
            "pnj_example": text(row.get("Ví dụ PNJ")) or None,
            "source_url": text(row.get("URL nguồn")) or None,
            "verification_status": text(row.get("Xác minh L3")) or None,
            "pnj_status": text(row.get("Trạng thái PNJ")) or None,
            "priority": text(row.get("Ưu tiên")) or None,
            "metrics_available": normalized_label(row.get("Metrics available")) == "y",
            "source_location": text(row.get("Source location")) or None,
        })
    apqc_codes = {row["hierarchy_id"] for row in apqc_rows}
    if len(apqc_rows) != 362 or len(apqc_codes) != 362:
        raise ValueError(f"Expected 362 unique APQC L3 rows, got {len(apqc_rows)} rows/{len(apqc_codes)} codes")

    source_documents: list[dict] = []
    invalid_classifications = Counter()
    for row_number, row in row_dicts(workbook[document_sheet_name], 2, 3):
        code = text(row.get("Mã tài liệu"))
        title = text(row.get("Tên tài liệu"))
        if not code or not title:
            continue
        raw_l3 = text(row.get("APQC L3 (code)"))
        l3_code = raw_l3 if raw_l3 in apqc_codes else None
        if raw_l3 and raw_l3 not in {"#N/A", "N/A"} and raw_l3 not in apqc_codes:
            invalid_classifications[raw_l3] += 1
        type_code, group_code = TYPE_MAP.get(normalized_label(row.get("Loại tài liệu")), ("other", "other"))
        source_documents.append({
            "source_row": row_number,
            "code": code,
            "title": title,
            "valid_from": iso_date(row.get("Ngày hiệu lực")),
            "type_code": type_code,
            "group_code": group_code,
            "issuing_department": text(row.get("Đơn vị soạn thảo")) or text(row.get("Phòng soạn thảo")) or None,
            "document_abstract": text(row.get("Mục đích kinh doanh")) or None,
            "applicability_note": text(row.get("Phạm vi")) or None,
            "l3_code": l3_code,
            "raw_l3": raw_l3 or None,
            "metadata": {
                "source_name": SOURCE_NAME,
                "source_row": row_number,
                "drafting_department": text(row.get("Phòng soạn thảo")) or None,
                "drafting_unit": text(row.get("Đơn vị soạn thảo")) or None,
                "attachment_type": text(row.get("Loại đính kèm")) or None,
                "final_sheet_updated": text(row.get("Đã update Sheet Final")) or None,
                "l0_review": text(row.get("L0 review")) or None,
                "pic": text(row.get("PIC")) or None,
                "apqc_version_1": text(row.get("APQC version 1")) or None,
                "source_concern": text(row.get("Concern từ lần 1 (để lại để nhớ thôi, làm xong thì xóa)")) or None,
                "apqc_l3_raw": raw_l3 or None,
                "apqc_l3_name": text(row.get("L3 (công thức)")) or None,
                "classification_justification": text(row.get("Justification")) or None,
                "related_l3": text(row.get("Related L3")) or None,
                "needs_calibration": text(row.get("Need calibration?")) or None,
                "business_purpose": text(row.get("Mục đích kinh doanh")) or None,
                "scope": text(row.get("Phạm vi")) or None,
                "main_flow": text(row.get("Luồng nghiệp vụ chính")) or None,
                "critical_decisions": text(row.get("Quyết định trọng yếu")) or None,
                "accountability_impact": text(row.get("Trách nhiệm & Quy mô tác động")) or None,
            },
        })
    if len(source_documents) != 1844:
        raise ValueError(f"Expected 1844 valid document rows, got {len(source_documents)}")

    versions_by_code: dict[str, list[dict]] = defaultdict(list)
    for row in source_documents:
        versions_by_code[row["code"]].append(row)
    for code, rows in versions_by_code.items():
        rows.sort(key=lambda item: (item["valid_from"] or "0000-00-00", item["source_row"]))
        document_id = str(uuid.uuid5(NAMESPACE, f"document:{code}"))
        for version_number, row in enumerate(rows, start=1):
            row["document_id"] = document_id
            row["version_id"] = str(uuid.uuid5(uuid.UUID(document_id), f"version:{version_number}:{row['source_row']}"))
            row["version_number"] = version_number

    prepare_sql = f"""
-- Generated outside Git from workbook checksum {checksum}. Prior records remain available for audit and rollback.
update public.apqc_frameworks set is_active=false where is_active;

insert into public.apqc_frameworks(code,name,version,source_name,source_license_note,imported_at,is_active)
values (
  'PNJ-APQC-CAT-VER2', 'APQC catalogue with controlled PNJ interpretation',
  'Process Map workbook 2026-10-29', 'APQC_CAT_ver2_Chi tiết Level 3',
  'Private controlled workbook; APQC wording remains source-attributed.', now(), true
)
on conflict (code) do update set
  name=excluded.name, version=excluded.version, source_name=excluded.source_name,
  source_license_note=excluded.source_license_note, imported_at=now(), is_active=true;

insert into public.apqc_elements(framework_id,pcf_id,hierarchy_id,level_number,parent_id,name,description,display_order)
select target.id, source.pcf_id, source.hierarchy_id, 1, null, source.name, source.description, source.display_order
from public.apqc_elements source
join public.apqc_frameworks source_framework on source_framework.id=source.framework_id and source_framework.code='APQC-PCF-CROSS-INDUSTRY'
join public.apqc_frameworks target on target.code='PNJ-APQC-CAT-VER2'
where source.level_number=1
on conflict (framework_id,hierarchy_id) do update set
  pcf_id=excluded.pcf_id,name=excluded.name,description=excluded.description,display_order=excluded.display_order;

insert into public.apqc_elements(framework_id,pcf_id,hierarchy_id,level_number,parent_id,name,description,display_order)
select target.id, source.pcf_id, source.hierarchy_id, 2, target_parent.id, source.name, source.description, source.display_order
from public.apqc_elements source
join public.apqc_frameworks source_framework on source_framework.id=source.framework_id and source_framework.code='APQC-PCF-CROSS-INDUSTRY'
join public.apqc_frameworks target on target.code='PNJ-APQC-CAT-VER2'
join public.apqc_elements source_parent on source_parent.id=source.parent_id
join public.apqc_elements target_parent on target_parent.framework_id=target.id and target_parent.hierarchy_id=source_parent.hierarchy_id
where source.level_number=2
on conflict (framework_id,hierarchy_id) do update set
  pcf_id=excluded.pcf_id,parent_id=excluded.parent_id,name=excluded.name,
  description=excluded.description,display_order=excluded.display_order;

update public.documents set archived_at=coalesce(archived_at,now()), updated_at=now()
where archived_at is null and document_code ~ '^TG-DOC-(00[1-9]|01[0-2])$';
update public.document_versions version set replaced_at=coalesce(version.replaced_at,now())
from public.documents document
where version.document_id=document.id and document.document_code ~ '^TG-DOC-(00[1-9]|01[0-2])$'
  and version.replaced_at is null;
"""
    write(args.output / "000_prepare.sql", prepare_sql)

    for batch_number, start in enumerate(range(0, len(apqc_rows), 20), start=1):
        batch = apqc_rows[start:start + 20]
        tag = f"apqc{batch_number}_{checksum[:8]}"
        sql = f"""
with source as (
  select * from jsonb_to_recordset({sql_json(batch, tag)}) as row(
    source_row integer, display_order integer, category_code text, parent_code text,
    hierarchy_id text, pcf_id text, name text, description text, pnj_description text,
    trigger_text text, output_text text, select_when text, do_not_select_when text,
    pnj_example text, source_url text, verification_status text, pnj_status text,
    priority text, metrics_available boolean, source_location text
  )
)
insert into public.apqc_elements(
  framework_id, pcf_id, hierarchy_id, level_number, parent_id, name, description, display_order,
  pnj_description, trigger_text, output_text, select_when, do_not_select_when, pnj_example,
  source_url, verification_status, pnj_status, priority, metrics_available, source_location
)
select framework.id, source.pcf_id, source.hierarchy_id, 3, parent.id, source.name, source.description,
       200000 + source.display_order, source.pnj_description, source.trigger_text, source.output_text,
       source.select_when, source.do_not_select_when, source.pnj_example, source.source_url,
       source.verification_status, source.pnj_status, source.priority, source.metrics_available, source.source_location
from source
join public.apqc_frameworks framework on framework.code = 'PNJ-APQC-CAT-VER2'
join public.apqc_elements parent on parent.framework_id = framework.id
  and parent.hierarchy_id = source.parent_code and parent.level_number = 2
on conflict (framework_id, hierarchy_id) do update set
  pcf_id=excluded.pcf_id, parent_id=excluded.parent_id, name=excluded.name, description=excluded.description,
  display_order=excluded.display_order, pnj_description=excluded.pnj_description,
  trigger_text=excluded.trigger_text, output_text=excluded.output_text, select_when=excluded.select_when,
  do_not_select_when=excluded.do_not_select_when, pnj_example=excluded.pnj_example,
  source_url=excluded.source_url, verification_status=excluded.verification_status,
  pnj_status=excluded.pnj_status, priority=excluded.priority,
  metrics_available=excluded.metrics_available, source_location=excluded.source_location;
"""
        write(args.output / f"1{batch_number:02d}_apqc.sql", sql)

    source_documents.sort(key=lambda item: (item["code"], item["version_number"]))

    document_batches = []
    for batch_number, start in enumerate(range(0, len(source_documents), args.document_batch_size), start=1):
        batch = source_documents[start:start + args.document_batch_size]
        document_batches.append(len(batch))
        tag = f"docs{batch_number}_{checksum[:8]}"
        sql = f"""
with source as (
  select * from jsonb_to_recordset({sql_json(batch, tag)}) as row(
    source_row integer, code text, title text, valid_from date, type_code text, group_code text,
    issuing_department text, document_abstract text, applicability_note text, l3_code text,
    raw_l3 text, metadata jsonb, document_id uuid, version_id uuid, version_number integer
  )
), inserted_documents as (
  insert into public.documents(document_code)
  select distinct source.code from source
  on conflict (document_code) do update set
    archived_at=null,
    updated_at=now()
  returning id, document_code
), prepared_source as (
  select source.*, document.id as target_document_id,
         legacy.base_version + source.version_number as target_version_number
  from source
  join inserted_documents document on document.document_code=source.code
  cross join lateral (
    select coalesce(max(existing.version_number),0) as base_version
    from public.document_versions existing
    where existing.document_id=document.id
      and existing.source_metadata->>'source_name' is distinct from {sql_text(SOURCE_NAME)}
  ) legacy
), inserted_versions as (
  insert into public.document_versions(
    id, document_id, version_number, status, title_vi, document_group, document_type_value_id,
    language, document_abstract, access_level, change_summary, valid_from,
    issuing_department, applicability_note, source_metadata, replaced_at
  )
  select source.version_id, source.target_document_id, source.target_version_number, 'draft', source.title,
         source.group_code, type_value.id, 'vi', source.document_abstract, 'internal',
         'Nhập có kiểm soát từ ' || {sql_text(SOURCE_NAME)} || '; dòng ' || source.source_row::text || '.',
         source.valid_from, source.issuing_department, source.applicability_note, source.metadata, null
  from prepared_source source
  join public.vocabularies vocabulary on vocabulary.code='document_type'
  join public.vocabulary_values type_value on type_value.vocabulary_id=vocabulary.id and type_value.code=source.type_code
  on conflict (document_id, version_number) do update set
    status=excluded.status, title_vi=excluded.title_vi,
    document_group=excluded.document_group, document_type_value_id=excluded.document_type_value_id,
    document_abstract=excluded.document_abstract, valid_from=excluded.valid_from,
    issuing_department=excluded.issuing_department, applicability_note=excluded.applicability_note,
    source_metadata=excluded.source_metadata, replaced_at=null, updated_at=now()
  returning id, document_id, version_number
)
insert into public.document_apqc_mappings(document_version_id, apqc_element_id, source_name, source_row)
select imported_version.id, element.id, {sql_text(SOURCE_NAME)}, source.source_row
from prepared_source source
join inserted_versions imported_version on imported_version.document_id=source.target_document_id and imported_version.version_number=source.target_version_number
join public.apqc_elements element on element.hierarchy_id=source.l3_code and element.level_number=3
join public.apqc_frameworks framework on framework.id=element.framework_id and framework.is_active
where source.l3_code is not null
on conflict (document_version_id, apqc_element_id) do update set
  source_name=excluded.source_name, source_row=excluded.source_row;
"""
        write(args.output / f"2{batch_number:02d}_documents.sql", sql)

    finish_sql = f"""
do $$
declare
  actual_apqc integer;
  actual_documents integer;
  actual_versions integer;
  actual_mappings integer;
begin
  select count(*) into actual_apqc from public.apqc_elements element join public.apqc_frameworks framework on framework.id=element.framework_id where framework.is_active and element.level_number=3;
  select count(*) into actual_documents from public.documents where archived_at is null;
  select count(*) into actual_versions from public.document_versions where replaced_at is null and source_metadata->>'source_name'={sql_text(SOURCE_NAME)};
  select count(*) into actual_mappings
  from public.document_apqc_mappings mapping
  join public.apqc_elements element on element.id=mapping.apqc_element_id
  join public.apqc_frameworks framework on framework.id=element.framework_id and framework.is_active
  where mapping.source_name={sql_text(SOURCE_NAME)};
  if actual_apqc <> 362 or actual_documents <> {len(versions_by_code)} or actual_versions <> {len(source_documents)} or actual_mappings <> {sum(1 for row in source_documents if row['l3_code'])} then
    raise exception 'Workbook import count mismatch: APQC %, documents %, versions %, mappings %', actual_apqc, actual_documents, actual_versions, actual_mappings;
  end if;
end $$;

insert into public.workbook_imports(source_name, source_sha256, apqc_row_count, document_row_count, notes)
values ({sql_text(SOURCE_NAME)}, '{checksum}', 362, {len(source_documents)},
  '1174 raw classifications; {sum(1 for row in source_documents if row['l3_code'])} valid APQC L3 links; invalid values remain in source_metadata for review.')
on conflict (source_sha256) do update set imported_at=now(), notes=excluded.notes;
"""
    write(args.output / "999_finish.sql", finish_sql)

    manifest = {
        "source_sha256": checksum,
        "apqc_sheet": apqc_sheet_name,
        "document_sheet": document_sheet_name,
        "apqc_rows": len(apqc_rows),
        "documents": len(versions_by_code),
        "document_versions": len(source_documents),
        "valid_apqc_links": sum(1 for row in source_documents if row["l3_code"]),
        "invalid_classifications": dict(invalid_classifications),
        "document_batches": document_batches,
        "generated_files": sorted(path.name for path in args.output.glob("*.sql")),
    }
    write(args.output / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
